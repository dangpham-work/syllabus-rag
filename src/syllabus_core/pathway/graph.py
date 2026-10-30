"""
graph.py
Xây dựng đồ thị tiên quyết môn học từ JSON đề cương và DanhMuc.xlsx (M4).

Dùng:
    from syllabus_core.pathway.graph import PrerequisiteGraph
    g = PrerequisiteGraph()
    g.load_from_syllabi("data/processed/syllabi")
"""

from __future__ import annotations

import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# CourseNode
# ---------------------------------------------------------------------------

@dataclass
class CourseNode:
    """Đại diện một môn học trong đồ thị."""
    code: str
    name: str
    credits: int = 3
    avg_bloom_level: float = 2.0
    recommended_semester: int = 1  # học kỳ khuyến nghị (từ DanhMuc.xlsx nếu có)

    def __hash__(self):
        return hash(self.code)

    def __eq__(self, other):
        return isinstance(other, CourseNode) and self.code == other.code


# ---------------------------------------------------------------------------
# PrerequisiteGraph
# ---------------------------------------------------------------------------

class PrerequisiteGraph:
    """
    Đồ thị có hướng không chu trình (DAG) biểu diễn quan hệ tiên quyết.

    Cạnh A → B nghĩa là A là tiên quyết của B.
    """

    def __init__(self):
        self._nodes: dict[str, CourseNode] = {}
        # _edges[A] = {B, C} nghĩa là A → B, A → C (A tiên quyết của B, C)
        self._edges: dict[str, set[str]] = {}
        # _reverse[B] = {A} nghĩa là A tiên quyết của B
        self._reverse: dict[str, set[str]] = {}
        # Theo dõi nguồn của mỗi cạnh: {(A,B): ['m1'|'xlsx']}
        self._edge_sources: dict[tuple[str, str], list[str]] = {}

    # ------------------------------------------------------------------
    # Thêm node / cạnh
    # ------------------------------------------------------------------

    def add_course(self, node: CourseNode) -> None:
        if node.code not in self._nodes:
            self._nodes[node.code] = node
            self._edges[node.code] = set()
            self._reverse[node.code] = set()

    def add_edge(self, from_code: str, to_code: str, source: str = "m1") -> None:
        """Thêm cạnh from_code → to_code (from là tiên quyết của to)."""
        for code in (from_code, to_code):
            if code not in self._nodes:
                self._nodes[code] = CourseNode(code=code, name=code)
                self._edges[code] = set()
                self._reverse[code] = set()
        self._edges[from_code].add(to_code)
        self._reverse[to_code].add(from_code)
        key = (from_code, to_code)
        if key not in self._edge_sources:
            self._edge_sources[key] = []
        if source not in self._edge_sources[key]:
            self._edge_sources[key].append(source)

    # ------------------------------------------------------------------
    # Load từ JSON đề cương
    # ------------------------------------------------------------------

    def load_from_syllabi(self, syllabi_dir: str | Path) -> None:
        """Load tất cả JSON đề cương và xây đồ thị tiên quyết."""
        syllabi_dir = Path(syllabi_dir)
        json_files = sorted(syllabi_dir.glob("*.json"))

        if not json_files:
            logger.warning("Không có JSON trong %s. Dùng synthetic graph.", syllabi_dir)
            self._load_synthetic()
            return

        for jf in json_files:
            try:
                with jf.open(encoding="utf-8") as f:
                    syl = json.load(f)
                code = syl.get("course_code") or jf.stem
                name = syl.get("course_name", code)
                credits = syl.get("credits", 3)

                # Tính avg_bloom từ CLOs
                clo_levels = [c.get("bloom_level", 2) for c in syl.get("clos", []) if c.get("bloom_level")]
                avg_bloom = sum(clo_levels) / len(clo_levels) if clo_levels else 2.0

                node = CourseNode(
                    code=code, name=name, credits=credits, avg_bloom_level=avg_bloom,
                )
                self.add_course(node)

                # Cạnh tiên quyết
                for prereq in syl.get("prerequisites", []):
                    if prereq and isinstance(prereq, str):
                        self.add_edge(prereq, code, source="m1")

            except Exception as e:
                logger.error("Lỗi đọc %s: %s", jf.name, e)

        logger.info(
            "Graph: %d môn, %d cạnh tiên quyết",
            len(self._nodes), sum(len(v) for v in self._edges.values()),
        )

    def _load_synthetic(self) -> None:
        """Graph tổng hợp khi không có dữ liệu thực."""
        sample = [
            CourseNode("841020", "Cơ sở lập trình", 3, 2.5, 1),
            CourseNode("841401", "Giải tích 1", 3, 2.0, 1),
            CourseNode("841108", "Cấu trúc dữ liệu và giải thuật", 3, 3.0, 2),
            CourseNode("841047", "Công nghệ phần mềm", 4, 3.5, 4),
            CourseNode("841408", "Kiểm thử phần mềm", 3, 3.8, 5),
            CourseNode("841470", "Tương tác người-máy", 3, 3.2, 6),
            CourseNode("841303", "Mạng máy tính", 3, 3.0, 3),
            CourseNode("841058", "Hệ điều hành mã nguồn mở", 3, 3.0, 3),
        ]
        edges = [
            ("841020", "841108"),
            ("841108", "841047"),
            ("841047", "841408"),
            ("841047", "841470"),
            ("841020", "841303"),
            ("841020", "841058"),
        ]
        for node in sample:
            self.add_course(node)
        for a, b in edges:
            self.add_edge(a, b, source="synthetic")
        logger.info("Synthetic graph: %d môn, %d cạnh", len(sample), len(edges))

    # ------------------------------------------------------------------
    # Load từ DanhMuc.xlsx
    # ------------------------------------------------------------------

    def load_from_xlsx(self, xlsx_path: str | Path) -> None:
        """Load quan hệ tiên quyết từ DanhMuc.xlsx (nếu có)."""
        xlsx_path = Path(xlsx_path)
        if not xlsx_path.exists():
            logger.warning("Không tìm thấy %s. Bỏ qua.", xlsx_path)
            return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path, read_only=True)
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows())]

            def col(name):
                for k in headers:
                    if name in k:
                        return headers.index(k)
                return -1

            code_col = col("mã")
            name_col = col("tên")
            sem_col  = col("học kỳ")
            prereq_col = col("tiên quyết")
            credit_col = col("tín chỉ")

            for row in ws.iter_rows(min_row=2, values_only=True):
                code = str(row[code_col]).strip() if code_col >= 0 and row[code_col] else None
                if not code or not code.startswith("84"):
                    continue
                name = str(row[name_col]).strip() if name_col >= 0 and row[name_col] else code
                credits = int(row[credit_col]) if credit_col >= 0 and row[credit_col] else 3
                sem = int(row[sem_col]) if sem_col >= 0 and row[sem_col] else 1

                if code in self._nodes:
                    self._nodes[code].recommended_semester = sem
                else:
                    self.add_course(CourseNode(code=code, name=name, credits=credits, recommended_semester=sem))

                prereq_val = row[prereq_col] if prereq_col >= 0 else None
                if prereq_val and str(prereq_val).strip().lower() not in ("", "không", "none"):
                    for prereq in str(prereq_val).split(","):
                        prereq = prereq.strip()
                        if prereq:
                            self.add_edge(prereq, code, source="xlsx")

            logger.info("Đã load %s vào graph.", xlsx_path.name)
        except Exception as e:
            logger.error("Lỗi đọc xlsx: %s", e)

    # ------------------------------------------------------------------
    # Kiểm tra
    # ------------------------------------------------------------------

    def check_cycles(self) -> list[list[str]]:
        """Tìm chu trình trong đồ thị (DFS). Returns list of cycles."""
        visited = set()
        rec_stack = set()
        cycles = []

        def dfs(node, path):
            visited.add(node)
            rec_stack.add(node)
            for neighbor in self._edges.get(node, set()):
                if neighbor not in visited:
                    dfs(neighbor, path + [neighbor])
                elif neighbor in rec_stack:
                    # Tìm thấy chu trình
                    cycle_start = path.index(neighbor)
                    cycles.append(path[cycle_start:] + [neighbor])
            rec_stack.discard(node)

        for node in self._nodes:
            if node not in visited:
                dfs(node, [node])

        return cycles

    def get_discrepancies(self) -> list[dict]:
        """Tìm cạnh chỉ có trong một nguồn (m1 vs xlsx)."""
        discrepancies = []
        for (a, b), sources in self._edge_sources.items():
            if "m1" in sources and "xlsx" not in sources:
                discrepancies.append({"from": a, "to": b, "only_in": "m1"})
            elif "xlsx" in sources and "m1" not in sources:
                discrepancies.append({"from": a, "to": b, "only_in": "xlsx"})
        return discrepancies

    def get_ancestors(self, course_code: str) -> set[str]:
        """Lấy tất cả tổ tiên (tiên quyết bắc cầu) của một môn."""
        ancestors = set()
        queue = list(self._reverse.get(course_code, set()))
        while queue:
            node = queue.pop()
            if node not in ancestors:
                ancestors.add(node)
                queue.extend(self._reverse.get(node, set()))
        return ancestors

    # ------------------------------------------------------------------
    # Properties & serialization
    # ------------------------------------------------------------------

    @property
    def nodes(self) -> dict[str, CourseNode]:
        return self._nodes

    @property
    def edges(self) -> dict[str, set[str]]:
        return self._edges

    def to_dict(self) -> dict:
        return {
            "nodes": {
                code: {
                    "name": n.name, "credits": n.credits,
                    "avg_bloom_level": n.avg_bloom_level,
                    "recommended_semester": n.recommended_semester,
                }
                for code, n in self._nodes.items()
            },
            "edges": {a: sorted(b) for a, b in self._edges.items()},
            "discrepancies": self.get_discrepancies(),
        }
